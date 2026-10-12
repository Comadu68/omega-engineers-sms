import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_card_borders(cell, border_color="BAE6FD", bg_color="F0F9FF"):
    set_cell_background(cell, bg_color)
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="12" w:space="0" w:color="{border_color}"/>
            <w:left w:val="single" w:sz="12" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="single" w:sz="12" w:space="0" w:color="{border_color}"/>
            <w:right w:val="single" w:sz="12" w:space="0" w:color="{border_color}"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="180" w:type="dxa"/>
            <w:bottom w:w="180" w:type="dxa"/>
            <w:left w:w="220" w:type="dxa"/>
            <w:right w:w="220" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def generate_ai_cards_docx():
    doc = Document()

    for s in doc.sections:
        s.top_margin = Inches(0.6)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.6)
        s.right_margin = Inches(0.6)

    primary_color = RGBColor(30, 58, 138)     # Navy #1E3A8A
    slate_dark = RGBColor(15, 23, 42)        # Slate 900
    slate_sub = RGBColor(71, 85, 105)        # Slate 600
    accent_blue = RGBColor(2, 132, 199)      # Sky blue #0284C7
    accent_purple = RGBColor(124, 58, 237)   # Purple #7C3AED
    accent_green = RGBColor(22, 101, 52)     # Green
    accent_orange = RGBColor(194, 65, 12)    # Orange

    # Sub-header tag
    p_tag = doc.add_paragraph()
    p_tag.paragraph_format.space_before = Pt(0)
    p_tag.paragraph_format.space_after = Pt(2)
    r_tag = p_tag.add_run("18  ACADEMIC INTEGRITY & DISCLOSURE")
    r_tag.font.name = "Calibri"
    r_tag.font.size = Pt(9.5)
    r_tag.font.bold = True
    r_tag.font.color.rgb = RGBColor(22, 163, 74) # Emerald green

    # Main Heading
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(14)
    r_title = p_title.add_run("Appendix: Generative AI Usage Declaration")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = primary_color

    # Intro
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_after = Pt(12)
    r_in = p_intro.add_run("The following declaration reflects the authentic distribution of Generative AI assistance utilized during the development lifecycle of the Omega Engineers Service Station Management System:")
    r_in.font.name = "Calibri"
    r_in.font.size = Pt(10)
    r_in.font.italic = True
    r_in.font.color.rgb = slate_sub

    # Table: 2x2 Grid with a spacing column or standard 2 columns
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_w = Inches(3.55)
    table.columns[0].width = col_w
    table.columns[1].width = col_w

    # Row 1, Col 0: Antigravity
    c00 = table.cell(0, 0)
    c00.width = col_w
    set_card_borders(c00, border_color="93C5FD", bg_color="F0F9FF") # Light blue card

    p_c00_title = c00.paragraphs[0]
    p_c00_title.paragraph_format.space_before = Pt(0)
    p_c00_title.paragraph_format.space_after = Pt(1)
    r_c00_icon = p_c00_title.add_run("⚡ Google DeepMind Antigravity")
    r_c00_icon.font.name = "Calibri"
    r_c00_icon.font.bold = True
    r_c00_icon.font.size = Pt(12)
    r_c00_icon.font.color.rgb = RGBColor(29, 78, 216)

    p_c00_sub = c00.add_paragraph()
    p_c00_sub.paragraph_format.space_before = Pt(0)
    p_c00_sub.paragraph_format.space_after = Pt(6)
    r_c00_sub = p_c00_sub.add_run("Full-Stack Coding & Prototype Implementation")
    r_c00_sub.font.name = "Calibri"
    r_c00_sub.font.size = Pt(9.5)
    r_c00_sub.font.bold = True
    r_c00_sub.font.color.rgb = slate_sub

    p_c00_body = c00.add_paragraph()
    p_c00_body.paragraph_format.space_before = Pt(0)
    p_c00_body.paragraph_format.space_after = Pt(4)
    r_b1_lbl = p_c00_body.add_run("Generated Output: ")
    r_b1_lbl.font.name = "Calibri"
    r_b1_lbl.font.bold = True
    r_b1_lbl.font.size = Pt(9)
    r_b1_lbl.font.color.rgb = slate_dark
    r_b1_txt = p_c00_body.add_run("Assisted in generating the full-stack Electron desktop codebase, modular IPC handlers, preload contextBridge, 13 domain services, React UI views, and the automated integration test suite (integration.test.cjs).")
    r_b1_txt.font.name = "Calibri"
    r_b1_txt.font.size = Pt(8.5)
    r_b1_txt.font.color.rgb = slate_dark

    p_c00_ver = c00.add_paragraph()
    p_c00_ver.paragraph_format.space_before = Pt(0)
    p_c00_ver.paragraph_format.space_after = Pt(0)
    r_v1_lbl = p_c00_ver.add_run("Team Verification: ")
    r_v1_lbl.font.name = "Calibri"
    r_v1_lbl.font.bold = True
    r_v1_lbl.font.size = Pt(9)
    r_v1_lbl.font.color.rgb = RGBColor(29, 78, 216)
    r_v1_txt = p_c00_ver.add_run("Every component was locally tested, verified via integration suites, checked for contextIsolation security, and validated with SQLite WAL mode execution.")
    r_v1_txt.font.name = "Calibri"
    r_v1_txt.font.size = Pt(8.5)
    r_v1_txt.font.color.rgb = slate_dark

    # Row 1, Col 1: Claude AI
    c01 = table.cell(0, 1)
    c01.width = col_w
    set_card_borders(c01, border_color="FED7AA", bg_color="FFF7ED") # Light Orange/Amber

    p_c01_title = c01.paragraphs[0]
    p_c01_title.paragraph_format.space_before = Pt(0)
    p_c01_title.paragraph_format.space_after = Pt(1)
    r_c01_icon = p_c01_title.add_run("✴ Claude AI (Anthropic)")
    r_c01_icon.font.name = "Calibri"
    r_c01_icon.font.bold = True
    r_c01_icon.font.size = Pt(12)
    r_c01_icon.font.color.rgb = RGBColor(194, 65, 12)

    p_c01_sub = c01.add_paragraph()
    p_c01_sub.paragraph_format.space_before = Pt(0)
    p_c01_sub.paragraph_format.space_after = Pt(6)
    r_c01_sub = p_c01_sub.add_run("Wireframing, Requirements & MoSCoW Prioritization")
    r_c01_sub.font.name = "Calibri"
    r_c01_sub.font.size = Pt(9.5)
    r_c01_sub.font.bold = True
    r_c01_sub.font.color.rgb = slate_sub

    p_c01_body = c01.add_paragraph()
    p_c01_body.paragraph_format.space_before = Pt(0)
    p_c01_body.paragraph_format.space_after = Pt(4)
    r_b2_lbl = p_c01_body.add_run("Generated Output: ")
    r_b2_lbl.font.name = "Calibri"
    r_b2_lbl.font.bold = True
    r_b2_lbl.font.size = Pt(9)
    r_b2_lbl.font.color.rgb = slate_dark
    r_b2_txt = p_c01_body.add_run("Assisted in creating interactive HTML wireframe layouts (Original_Wireframes.html), formalizing Functional & Non-Functional specifications, and structuring the MoSCoW prioritization matrix.")
    r_b2_txt.font.name = "Calibri"
    r_b2_txt.font.size = Pt(8.5)
    r_b2_txt.font.color.rgb = slate_dark

    p_c01_ver = c01.add_paragraph()
    p_c01_ver.paragraph_format.space_before = Pt(0)
    p_c01_ver.paragraph_format.space_after = Pt(0)
    r_v2_lbl = p_c01_ver.add_run("Team Verification: ")
    r_v2_lbl.font.name = "Calibri"
    r_v2_lbl.font.bold = True
    r_v2_lbl.font.size = Pt(9)
    r_v2_lbl.font.color.rgb = RGBColor(194, 65, 12)
    r_v2_txt = p_c01_ver.add_run("Wireframes and prioritized requirements were cross-verified against actual counter workflows and operational boundaries at Omega Engineers.")
    r_v2_txt.font.name = "Calibri"
    r_v2_txt.font.size = Pt(8.5)
    r_v2_txt.font.color.rgb = slate_dark

    # Row 2, Col 0: Gemini
    c10 = table.cell(1, 0)
    c10.width = col_w
    set_card_borders(c10, border_color="E9D5FF", bg_color="FAF5FF") # Light Purple

    p_c10_title = c10.paragraphs[0]
    p_c10_title.paragraph_format.space_before = Pt(0)
    p_c10_title.paragraph_format.space_after = Pt(1)
    r_c10_icon = p_c10_title.add_run("✦ Gemini (Google)")
    r_c10_icon.font.name = "Calibri"
    r_c10_icon.font.bold = True
    r_c10_icon.font.size = Pt(12)
    r_c10_icon.font.color.rgb = RGBColor(126, 34, 206)

    p_c10_sub = c10.add_paragraph()
    p_c10_sub.paragraph_format.space_before = Pt(0)
    p_c10_sub.paragraph_format.space_after = Pt(6)
    r_c10_sub = p_c10_sub.add_run("Idea Generation, Technical Research & Domain Q&A")
    r_c10_sub.font.name = "Calibri"
    r_c10_sub.font.size = Pt(9.5)
    r_c10_sub.font.bold = True
    r_c10_sub.font.color.rgb = slate_sub

    p_c10_body = c10.add_paragraph()
    p_c10_body.paragraph_format.space_before = Pt(0)
    p_c10_body.paragraph_format.space_after = Pt(4)
    r_b3_lbl = p_c10_body.add_run("Generated Output: ")
    r_b3_lbl.font.name = "Calibri"
    r_b3_lbl.font.bold = True
    r_b3_lbl.font.size = Pt(9)
    r_b3_lbl.font.color.rgb = slate_dark
    r_b3_txt = p_c10_body.add_run("Helped brainstorm system architectural alternatives, clarify unfamiliar technical concepts, explore service station business practices, and investigate FIFO inventory valuation models.")
    r_b3_txt.font.name = "Calibri"
    r_b3_txt.font.size = Pt(8.5)
    r_b3_txt.font.color.rgb = slate_dark

    p_c10_ver = c10.add_paragraph()
    p_c10_ver.paragraph_format.space_before = Pt(0)
    p_c10_ver.paragraph_format.space_after = Pt(0)
    r_v3_lbl = p_c10_ver.add_run("Team Verification: ")
    r_v3_lbl.font.name = "Calibri"
    r_v3_lbl.font.bold = True
    r_v3_lbl.font.size = Pt(9)
    r_v3_lbl.font.color.rgb = RGBColor(126, 34, 206)
    r_v3_txt = p_c10_ver.add_run("All conceptual suggestions and research findings were audited by the team against offline deployment constraints and official technical documentation.")
    r_v3_txt.font.name = "Calibri"
    r_v3_txt.font.size = Pt(8.5)
    r_v3_txt.font.color.rgb = slate_dark

    # Row 2, Col 1: Microsoft Copilot
    c11 = table.cell(1, 1)
    c11.width = col_w
    set_card_borders(c11, border_color="BBF7D0", bg_color="F0FDF4") # Light Green

    p_c11_title = c11.paragraphs[0]
    p_c11_title.paragraph_format.space_before = Pt(0)
    p_c11_title.paragraph_format.space_after = Pt(1)
    r_c11_icon = p_c11_title.add_run("❖ Microsoft Copilot")
    r_c11_icon.font.name = "Calibri"
    r_c11_icon.font.bold = True
    r_c11_icon.font.size = Pt(12)
    r_c11_icon.font.color.rgb = RGBColor(21, 128, 61)

    p_c11_sub = c11.add_paragraph()
    p_c11_sub.paragraph_format.space_before = Pt(0)
    p_c11_sub.paragraph_format.space_after = Pt(6)
    r_c11_sub = p_c11_sub.add_run("Documentation Authoring, Alignment & Grammar Polish")
    r_c11_sub.font.name = "Calibri"
    r_c11_sub.font.size = Pt(9.5)
    r_c11_sub.font.bold = True
    r_c11_sub.font.color.rgb = slate_sub

    p_c11_body = c11.add_paragraph()
    p_c11_body.paragraph_format.space_before = Pt(0)
    p_c11_body.paragraph_format.space_after = Pt(4)
    r_b4_lbl = p_c11_body.add_run("Generated Output: ")
    r_b4_lbl.font.name = "Calibri"
    r_b4_lbl.font.bold = True
    r_b4_lbl.font.size = Pt(9)
    r_b4_lbl.font.color.rgb = slate_dark
    r_b4_txt = p_c11_body.add_run("Assisted in drafting documentation sections, proofreading prose, correcting grammatical errors, aligning document layout, and checking formatting consistency across report appendices.")
    r_b4_txt.font.name = "Calibri"
    r_b4_txt.font.size = Pt(8.5)
    r_b4_txt.font.color.rgb = slate_dark

    p_c11_ver = c11.add_paragraph()
    p_c11_ver.paragraph_format.space_before = Pt(0)
    p_c11_ver.paragraph_format.space_after = Pt(0)
    r_v4_lbl = p_c11_ver.add_run("Team Verification: ")
    r_v4_lbl.font.name = "Calibri"
    r_v4_lbl.font.bold = True
    r_v4_lbl.font.size = Pt(9)
    r_v4_lbl.font.color.rgb = RGBColor(21, 128, 61)
    r_v4_txt = p_c11_ver.add_run("The team conducted manual peer-reviews on all drafted text to ensure academic tone, accurate technical claims, and strict compliance with university rubrics.")
    r_v4_txt.font.name = "Calibri"
    r_v4_txt.font.size = Pt(8.5)
    r_v4_txt.font.color.rgb = slate_dark

    # Bottom Spacing
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Footnote
    p_foot = doc.add_paragraph()
    p_foot.paragraph_format.space_before = Pt(10)
    p_foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_ft = p_foot.add_run("Page 18 • Academic Integrity & Generative AI Disclosure • Omega Engineers")
    r_ft.font.name = "Calibri"
    r_ft.font.size = Pt(8.5)
    r_ft.font.italic = True
    r_ft.font.color.rgb = slate_sub

    out_file = "D:/Projects/Omega Service/Generative_AI_Usage_Declaration_Cards.docx"
    doc.save(out_file)
    print(f"Saved: {out_file}")

if __name__ == "__main__":
    generate_ai_cards_docx()
