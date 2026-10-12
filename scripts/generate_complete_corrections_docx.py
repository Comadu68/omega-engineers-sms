import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_card_borders(cell, border_color="BAE6FD", bg_color="F0F9FF"):
    set_cell_background(cell, bg_color)
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="12" w:space="0" w:color="{border_color}"/>
            <w:left w:val="single" w:sz="12" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="single" w:sz="12" w:space="0" w:color="{border_color}"/>
            <w:right w:val="single" w:sz="12" w:space="0" w:color="{border_color}"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="160" w:type="dxa"/>
            <w:bottom w:w="160" w:type="dxa"/>
            <w:left w:w="200" w:type="dxa"/>
            <w:right w:w="200" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="none"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="none"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def generate_complete_corrections():
    doc = Document()

    for s in doc.sections:
        s.top_margin = Inches(0.65)
        s.bottom_margin = Inches(0.65)
        s.left_margin = Inches(0.65)
        s.right_margin = Inches(0.65)

    primary_color = RGBColor(30, 58, 138)     # Navy #1E3A8A
    secondary_color = RGBColor(71, 85, 105)  # Slate #475569
    dark_text = RGBColor(15, 23, 42)         # Slate 900 #0F172A

    # Document Header
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("Omega Engineers SSMS — Comprehensive Report Corrections")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = primary_color

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("Verified Real-World Alignment for Generative AI Declarations (Cards & Table), Git Profile, and Architecture Tree")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(10.5)
    r_sub.font.color.rgb = secondary_color
    r_sub.font.italic = True

    # ==========================================================
    # SECTION 1: PAGE 18 CARD-STYLE AI USAGE DECLARATION
    # ==========================================================
    h_cards = doc.add_paragraph()
    h_cards.paragraph_format.space_before = Pt(6)
    h_cards.paragraph_format.space_after = Pt(2)
    r_hcards = h_cards.add_run("1. Page 18: Appendix — Generative AI Usage Declaration (Card Layout)")
    r_hcards.font.name = "Calibri"
    r_hcards.font.size = Pt(13)
    r_hcards.font.bold = True
    r_hcards.font.color.rgb = primary_color

    p_tag = doc.add_paragraph()
    p_tag.paragraph_format.space_after = Pt(8)
    r_tag = p_tag.add_run("Corresponds to Slide / Page 18: Academic Integrity & Disclosure (Cards Grid)")
    r_tag.font.name = "Calibri"
    r_tag.font.size = Pt(9.5)
    r_tag.font.italic = True
    r_tag.font.color.rgb = secondary_color

    tbl_cards = doc.add_table(rows=2, cols=2)
    tbl_cards.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_cards.autofit = False

    col_w = Inches(3.55)
    tbl_cards.columns[0].width = col_w
    tbl_cards.columns[1].width = col_w

    # Card 0,0: Antigravity
    c00 = tbl_cards.cell(0, 0)
    c00.width = col_w
    set_card_borders(c00, border_color="93C5FD", bg_color="F0F9FF")
    p = c00.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("⚡ Google DeepMind Antigravity")
    r.font.name = "Calibri"; r.font.bold = True; r.font.size = Pt(11); r.font.color.rgb = RGBColor(29, 78, 216)
    p2 = c00.add_paragraph()
    p2.paragraph_format.space_after = Pt(4)
    r = p2.add_run("Full-Stack Coding & Prototype Implementation")
    r.font.name = "Calibri"; r.font.size = Pt(9); r.font.bold = True; r.font.color.rgb = secondary_color
    p3 = c00.add_paragraph()
    p3.paragraph_format.space_after = Pt(3)
    r = p3.add_run("Generated Output: "); r.font.name = "Calibri"; r.font.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = dark_text
    r = p3.add_run("Assisted in generating the full-stack Electron desktop codebase, modular IPC handlers, preload contextBridge, 13 domain services, React UI views, and the automated integration test suite (integration.test.cjs)."); r.font.name = "Calibri"; r.font.size = Pt(8); r.font.color.rgb = dark_text
    p4 = c00.add_paragraph()
    p4.paragraph_format.space_after = Pt(0)
    r = p4.add_run("Team Verification: "); r.font.name = "Calibri"; r.font.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = RGBColor(29, 78, 216)
    r = p4.add_run("Every component was locally tested, verified via integration suites, checked for contextIsolation security, and validated with SQLite WAL mode execution."); r.font.name = "Calibri"; r.font.size = Pt(8); r.font.color.rgb = dark_text

    # Card 0,1: Claude AI
    c01 = tbl_cards.cell(0, 1)
    c01.width = col_w
    set_card_borders(c01, border_color="FED7AA", bg_color="FFF7ED")
    p = c01.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("✴ Claude AI (Anthropic)")
    r.font.name = "Calibri"; r.font.bold = True; r.font.size = Pt(11); r.font.color.rgb = RGBColor(194, 65, 12)
    p2 = c01.add_paragraph()
    p2.paragraph_format.space_after = Pt(4)
    r = p2.add_run("Wireframing, Requirements & MoSCoW Prioritization")
    r.font.name = "Calibri"; r.font.size = Pt(9); r.font.bold = True; r.font.color.rgb = secondary_color
    p3 = c01.add_paragraph()
    p3.paragraph_format.space_after = Pt(3)
    r = p3.add_run("Generated Output: "); r.font.name = "Calibri"; r.font.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = dark_text
    r = p3.add_run("Assisted in creating interactive HTML wireframe layouts (Original_Wireframes.html), formalizing Functional & Non-Functional specifications, and structuring the MoSCoW prioritization matrix."); r.font.name = "Calibri"; r.font.size = Pt(8); r.font.color.rgb = dark_text
    p4 = c01.add_paragraph()
    p4.paragraph_format.space_after = Pt(0)
    r = p4.add_run("Team Verification: "); r.font.name = "Calibri"; r.font.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = RGBColor(194, 65, 12)
    r = p4.add_run("Wireframes and prioritized requirements were cross-verified against actual counter workflows and operational boundaries at Omega Engineers."); r.font.name = "Calibri"; r.font.size = Pt(8); r.font.color.rgb = dark_text

    # Card 1,0: Gemini
    c10 = tbl_cards.cell(1, 0)
    c10.width = col_w
    set_card_borders(c10, border_color="E9D5FF", bg_color="FAF5FF")
    p = c10.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("✦ Gemini (Google)")
    r.font.name = "Calibri"; r.font.bold = True; r.font.size = Pt(11); r.font.color.rgb = RGBColor(126, 34, 206)
    p2 = c10.add_paragraph()
    p2.paragraph_format.space_after = Pt(4)
    r = p2.add_run("Idea Generation, Technical Research & Domain Q&A")
    r.font.name = "Calibri"; r.font.size = Pt(9); r.font.bold = True; r.font.color.rgb = secondary_color
    p3 = c10.add_paragraph()
    p3.paragraph_format.space_after = Pt(3)
    r = p3.add_run("Generated Output: "); r.font.name = "Calibri"; r.font.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = dark_text
    r = p3.add_run("Helped brainstorm system architectural alternatives, clarify unfamiliar technical concepts, explore service station business practices, and investigate FIFO inventory valuation models."); r.font.name = "Calibri"; r.font.size = Pt(8); r.font.color.rgb = dark_text
    p4 = c10.add_paragraph()
    p4.paragraph_format.space_after = Pt(0)
    r = p4.add_run("Team Verification: "); r.font.name = "Calibri"; r.font.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = RGBColor(126, 34, 206)
    r = p4.add_run("All conceptual suggestions and research findings were audited by the team against offline deployment constraints and official technical documentation."); r.font.name = "Calibri"; r.font.size = Pt(8); r.font.color.rgb = dark_text

    # Card 1,1: Copilot
    c11 = tbl_cards.cell(1, 1)
    c11.width = col_w
    set_card_borders(c11, border_color="BBF7D0", bg_color="F0FDF4")
    p = c11.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("❖ Microsoft Copilot")
    r.font.name = "Calibri"; r.font.bold = True; r.font.size = Pt(11); r.font.color.rgb = RGBColor(21, 128, 61)
    p2 = c11.add_paragraph()
    p2.paragraph_format.space_after = Pt(4)
    r = p2.add_run("Documentation Authoring, Alignment & Grammar Polish")
    r.font.name = "Calibri"; r.font.size = Pt(9); r.font.bold = True; r.font.color.rgb = secondary_color
    p3 = c11.add_paragraph()
    p3.paragraph_format.space_after = Pt(3)
    r = p3.add_run("Generated Output: "); r.font.name = "Calibri"; r.font.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = dark_text
    r = p3.add_run("Assisted in drafting documentation sections, proofreading prose, correcting grammatical errors, aligning document layout, and checking formatting consistency across report appendices."); r.font.name = "Calibri"; r.font.size = Pt(8); r.font.color.rgb = dark_text
    p4 = c11.add_paragraph()
    p4.paragraph_format.space_after = Pt(0)
    r = p4.add_run("Team Verification: "); r.font.name = "Calibri"; r.font.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = RGBColor(21, 128, 61)
    r = p4.add_run("The team conducted manual peer-reviews on all drafted text to ensure academic tone, accurate technical claims, and strict compliance with university rubrics."); r.font.name = "Calibri"; r.font.size = Pt(8); r.font.color.rgb = dark_text

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # ==========================================================
    # SECTION 2: PAGE 45 DETAILED AI TOOLS TABLE
    # ==========================================================
    h_table = doc.add_paragraph()
    h_table.paragraph_format.space_before = Pt(8)
    h_table.paragraph_format.space_after = Pt(2)
    r_htable = h_table.add_run("2. Page 45: AI Tools Detailed Verification Table")
    r_htable.font.name = "Calibri"
    r_htable.font.size = Pt(13)
    r_htable.font.bold = True
    r_htable.font.color.rgb = primary_color

    tbl_ai = doc.add_table(rows=1, cols=4)
    tbl_ai.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_ai.autofit = False
    set_table_borders(tbl_ai, color="CBD5E1")

    widths_ai = [Inches(1.2), Inches(1.9), Inches(1.9), Inches(2.1)]
    for i, w in enumerate(widths_ai):
        tbl_ai.columns[i].width = w

    headers_ai = ["AI tools used", "Purpose of use", "Generated outputs", "How the outputs were verified and refined by the team"]
    hdr_cells_ai = tbl_ai.rows[0].cells
    for i, h_text in enumerate(headers_ai):
        hdr_cells_ai[i].width = widths_ai[i]
        set_cell_background(hdr_cells_ai[i], "1E3A8A")
        set_cell_margins(hdr_cells_ai[i], top=90, bottom=90, left=110, right=110)
        p = hdr_cells_ai[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h_text); r.font.name = "Calibri"; r.font.size = Pt(9); r.font.bold = True; r.font.color.rgb = RGBColor(255, 255, 255)

    data_ai = [
        (
            "Gemini\n(Google)",
            "Conceptual ideation, exploratory research on automotive service workflows, and technical inquiry into unfamiliar concepts and architecture options.",
            "Brainstormed system concepts, domain knowledge on service station operations, comparative research on offline database options, and FIFO inventory logic.",
            "The team reviewed research insights against actual counter operations at Omega Engineers, validating feasibility and filtering applicable ideas."
        ),
        (
            "Claude\n(Anthropic)",
            "Designing interactive UI wireframes, formalizing Software Requirements Specification (SRS), and structuring MoSCoW prioritization.",
            "Interactive HTML/CSS wireframe prototypes (Original_Wireframes.html), structured Functional & Non-Functional requirements list, and prioritized MoSCoW matrix.",
            "Business Developer H.M.P.S.D.W. Herath cross-verified each wireframe and requirement against physical station workflows, ensuring operational role boundaries were strictly preserved."
        ),
        (
            "Microsoft\nCopilot",
            "Academic documentation preparation, text alignment, formatting consistency, and grammatical enhancement across report sections.",
            "Polished academic prose, corrected grammatical phrasing, coherent technical section transitions, and formatted documentation tables.",
            "The entire team conducted peer reviews on all drafted sections to ensure academic rigor, accurate student details (724701588), and rubric compliance."
        ),
        (
            "Google DeepMind\nAntigravity",
            "Full-stack application code generation, interactive prototype development, modular Electron/React architecture implementation, and test automation.",
            "Production-ready Electron desktop source code (main process, IPC handlers, preload contextBridge, React UI pages, 13 domain services, and automated integration test suite).",
            "Backend Lead W.M.P.R. Weerasinghe and developers validated code via automated integration tests (tests/integration.test.cjs), verified contextIsolation security, and built the NSIS Windows installer."
        )
    ]

    for idx, (tool, purpose, outputs, verification) in enumerate(data_ai):
        row = tbl_ai.add_row()
        bg = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
        for i, val in enumerate([tool, purpose, outputs, verification]):
            c = row.cells[i]
            c.width = widths_ai[i]
            set_cell_background(c, bg)
            set_cell_margins(c, top=80, bottom=80, left=110, right=110)
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val); r.font.name = "Calibri"; r.font.size = Pt(8.5)
            if i == 0:
                r.font.bold = True; r.font.color.rgb = primary_color
            else:
                r.font.color.rgb = dark_text

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # ==========================================================
    # SECTION 3: APPENDIX D - REAL GIT METRICS & ARCHITECTURE
    # ==========================================================
    h_git = doc.add_paragraph()
    h_git.paragraph_format.space_before = Pt(8)
    h_git.paragraph_format.space_after = Pt(2)
    r_hgit = h_git.add_run("3. Appendix D: Real Git Repository Metrics & Modular Tree")
    r_hgit.font.name = "Calibri"
    r_hgit.font.size = Pt(13)
    r_hgit.font.bold = True
    r_hgit.font.color.rgb = primary_color

    profile_items = [
        ("• Repository URL: ", "https://github.com/Comadu68/omega-engineers-sms"),
        ("• Primary / Default Branch: ", "main"),
        ("• Branching Strategy: ", "Centralized Mainline (Trunk-Based) Milestone Delivery"),
        ("• Total Tracked Commits: ", "3 consolidated milestone releases (118 files, 20,200+ lines of production code)"),
        ("• Active Committer / Lead: ", "Comadu68 <pathum15333@gmail.com>"),
        ("• Quality Assurance & Verification Gates: ", "Automated Suite (tests/integration.test.cjs), ContextIsolation runtime validation, and Electron Builder NSIS release packaging tests.")
    ]

    for label, val in profile_items:
        p_item = doc.add_paragraph()
        p_item.paragraph_format.space_before = Pt(1)
        p_item.paragraph_format.space_after = Pt(2)
        p_item.paragraph_format.left_indent = Inches(0.2)
        r_lbl = p_item.add_run(label); r_lbl.font.name = "Calibri"; r_lbl.font.bold = True; r_lbl.font.size = Pt(9.5); r_lbl.font.color.rgb = dark_text
        r_v = p_item.add_run(val); r_v.font.name = "Calibri"; r_v.font.size = Pt(9.5)
        r_v.font.color.rgb = primary_color if "http" in val else dark_text
        if "http" in val: r_v.font.underline = True

    # Commits table
    tbl_commits = doc.add_table(rows=1, cols=4)
    tbl_commits.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_commits.autofit = False
    set_table_borders(tbl_commits, color="CBD5E1")

    widths_c = [Inches(0.9), Inches(1.1), Inches(1.7), Inches(3.4)]
    for i, w in enumerate(widths_c):
        tbl_commits.columns[i].width = w

    c_hdr = tbl_commits.rows[0].cells
    hdr_titles = ["Commit Hash", "Date", "Author", "Commit Message & Deliverable Scope"]
    for i, t in enumerate(hdr_titles):
        c_hdr[i].width = widths_c[i]
        set_cell_background(c_hdr[i], "1E3A8A")
        set_cell_margins(c_hdr[i], top=80, bottom=80, left=100, right=100)
        p = c_hdr[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
        r = p.add_run(t); r.font.name = "Calibri"; r.font.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = RGBColor(255, 255, 255)

    commits_data = [
        ("dc879cf", "2026-10-09", "Comadu68", "Update Readme.md\nAdded system deployment notes, execution prerequisites, and run guide."),
        ("8a80cdd", "2026-10-07", "Comadu68", "Initial commit: Omega Service Station System\nIntegrated full core application milestone: 118 files, 20,245 lines (Electron main/preload, SQLite schema & seed, 13 domain services, 13 React UI views, automated test suite)."),
        ("0090353", "2026-10-07", "Comadu68", "Initial commit\nInitialized Git repository and base workspace structure.")
    ]

    for idx, (chash, cdate, cauthor, cdesc) in enumerate(commits_data):
        row = tbl_commits.add_row()
        bg = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
        for i, val in enumerate([chash, cdate, cauthor, cdesc]):
            c = row.cells[i]
            c.width = widths_c[i]
            set_cell_background(c, bg)
            set_cell_margins(c, top=70, bottom=70, left=100, right=100)
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val); r.font.name = "Calibri"; r.font.size = Pt(8)
            if i == 0:
                r.font.name = "Consolas"; r.font.bold = True; r.font.color.rgb = primary_color
            else:
                r.font.color.rgb = dark_text

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Modular Tree
    tree_text = """omega-engineers-sms/
|-- /src/
|   |-- /main/                      # Electron Main Process (OS lifecycle, window orchestration)
|   |   |-- main.cjs                # Main entry point & browser window initialization
|   |   |-- /ipc/                   # Modular IPC channel registration handlers
|   |   |   |-- index.cjs           # Aggregator for all IPC registrations
|   |   |   |-- auth.ipc.cjs        # Authentication & user management IPC
|   |   |   |-- jobs.ipc.cjs        # Job card workflow & task execution IPC
|   |   |   |-- inventory.ipc.cjs   # Inventory stock, batches, & FIFO tracking IPC
|   |   |   |-- invoices.ipc.cjs    # Billing calculation & invoice generation IPC
|   |   |   |-- backup.ipc.cjs      # Database backup & restore IPC (.bak)
|   |   |   `-- ...                 # appointments, customers, mechanics, reports, etc.
|   |   `-- /services/              # Domain-Driven Business Logic & Security Services
|   |       |-- authService.cjs     # PBKDF2 hashing, per-user salt generation, RBAC
|   |       |-- jobService.cjs      # Job card state machine & mechanic assignments
|   |       |-- inventoryService.cjs# Stock thresholds, FIFO batch deductions, valuations
|   |       |-- invoiceService.cjs  # Parts + Labor + Tax billing engine & balance tracking
|   |       |-- backupService.cjs   # Automated SQLite database backup & validation
|   |       `-- ...                 # customerService, mechanicService, reportService, etc.
|   |
|   |-- /preload/                   # Secure IPC Context Isolation Layer
|   |   `-- preload.cjs             # contextBridge.exposeInMainWorld('omega', ...) API
|   |
|   |-- /database/                  # Embedded SQLite Persistence Layer
|   |   |-- db.cjs                  # better-sqlite3 connection with WAL mode & foreign keys
|   |   |-- schema.sql              # Normalized DDL schema (19 relational tables & indexes)
|   |   `-- seed.sql                # Default initial seed data & presets
|   |
|   `-- /renderer/                  # Presentation Layer (React.js + Vite UI)
|       |-- App.jsx                 # Application root & authentication router
|       |-- main.jsx                # React DOM entry point
|       |-- /components/            # Reusable UI tokens (Sidebar, Topbar, Modal, Badge)
|       |-- /pages/                 # Role-based screens & functional views
|       |   |-- Dashboard.jsx       # Operational metrics & status overview
|       |   |-- JobCards.jsx        # Job card management & dispatching
|       |   |-- Inventory.jsx       # Stock tracking, batch intake, supplier management
|       |   |-- Invoices.jsx        # Invoicing, payment collection, receipts
|       |   |-- InvoicePrintView.jsx# Print layout for A4 invoices & thermal print
|       |   |-- Reports.jsx         # Financial, job performance, and inventory reports
|       |   `-- ...                 # Appointments, Customers, Mechanics, Settings, Login
|       `-- /styles/                # Design system & CSS styling
|           `-- global.css          # Color tokens, typography, and @media print rules
|
`-- /tests/                         # Automated Verification Suite
    `-- integration.test.cjs        # End-to-end integration & database regression tests"""

    tree_tbl = doc.add_table(rows=1, cols=1)
    tree_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tree_tbl.autofit = False
    tree_tbl.columns[0].width = Inches(7.1)
    tree_cell = tree_tbl.cell(0, 0)
    set_cell_background(tree_cell, "0F172A")
    set_cell_margins(tree_cell, top=100, bottom=100, left=140, right=140)

    tcPr = tree_cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="12" w:space="0" w:color="334155"/>
            <w:top w:val="single" w:sz="12" w:space="0" w:color="334155"/>
            <w:right w:val="single" w:sz="12" w:space="0" w:color="334155"/>
            <w:bottom w:val="single" w:sz="12" w:space="0" w:color="334155"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)

    tp = tree_cell.paragraphs[0]
    tp.paragraph_format.space_before = Pt(0); tp.paragraph_format.space_after = Pt(0)
    tp.paragraph_format.line_spacing = Pt(11)

    r_tree = tp.add_run(tree_text)
    r_tree.font.name = "Consolas"; r_tree.font.size = Pt(7.5); r_tree.font.color.rgb = RGBColor(226, 232, 240)

    out_file = "D:/Projects/Omega Service/Omega_Service_Report_Corrections_Complete.docx"
    doc.save(out_file)
    print(f"Saved complete corrections: {out_file}")

if __name__ == "__main__":
    generate_complete_corrections()
