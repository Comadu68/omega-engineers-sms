import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="none"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="none"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def generate_ai_tools_docx():
    doc = Document()

    for s in doc.sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.7)
        s.left_margin = Inches(0.7)
        s.right_margin = Inches(0.7)

    primary_color = RGBColor(30, 58, 138)     # Navy #1E3A8A
    secondary_color = RGBColor(71, 85, 105)  # Slate #475569
    dark_text = RGBColor(15, 23, 42)         # Slate 900 #0F172A

    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("AI Tools Declaration Table (Corrected & Verified)")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(18)
    r_title.font.bold = True
    r_title.font.color.rgb = primary_color

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(12)
    r_sub = p_sub.add_run("Accurate Reflection of Project AI Usage for Academic Transparency & Evaluation")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(10.5)
    r_sub.font.color.rgb = secondary_color
    r_sub.font.italic = True

    # Note
    p_note = doc.add_paragraph()
    p_note.paragraph_format.space_after = Pt(12)
    r_n = p_note.add_run("In the previous version of the report, the tools' roles were inadvertently swapped (e.g. attributing code generation to Copilot and report writing to Antigravity). Below is the accurate, professionally phrased declaration matching the actual development workflows.")
    r_n.font.name = "Calibri"
    r_n.font.size = Pt(9.5)
    r_n.font.color.rgb = dark_text

    # Table
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="CBD5E1")

    widths = [Inches(1.2), Inches(1.9), Inches(1.9), Inches(2.1)]
    for i, w in enumerate(widths):
        table.columns[i].width = w

    headers = [
        "AI tools used",
        "Purpose of use",
        "Generated outputs",
        "How the outputs were verified and refined by the team"
    ]

    hdr_cells = table.rows[0].cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].width = widths[i]
        set_cell_background(hdr_cells[i], "1E3A8A")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h_text)
        r.font.name = "Calibri"
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    data = [
        (
            "Gemini\n(Google)",
            "Conceptual ideation, exploratory research on automotive service workflows, and technical inquiry into unfamiliar concepts and architecture options.",
            "Brainstormed system concepts, domain knowledge on service station operations, comparative research on offline database options, and FIFO inventory logic.",
            "The team reviewed research insights against actual counter operations at Omega Engineers, validating feasibility and filtering applicable ideas."
        ),
        (
            "Claude\n(Anthropic)",
            "Designing interactive UI wireframes, formalizing Software Requirements Specification (SRS), and structuring MoSCoW prioritization.",
            "Interactive HTML/CSS wireframe prototypes (Original_Wireframes.html), structured Functional & Non-Functional requirements list, and prioritized MoSCoW matrix.",
            "Business Developer H.M.P.S.D.W. Herath cross-verified each wireframe and requirement against physical station workflows, ensuring operational role boundaries were strictly preserved."
        ),
        (
            "Microsoft\nCopilot",
            "Academic documentation preparation, text alignment, formatting consistency, and grammatical enhancement across report sections.",
            "Polished academic prose, corrected grammatical phrasing, coherent technical section transitions, and formatted documentation tables.",
            "The entire team conducted peer reviews on all drafted sections to ensure academic rigor, accurate student details (724701588), and rubric compliance."
        ),
        (
            "Google DeepMind\nAntigravity",
            "Full-stack application code generation, interactive prototype development, modular Electron/React architecture implementation, and test automation.",
            "Production-ready Electron desktop source code (main process, IPC handlers, preload contextBridge, React UI pages, 13 domain services, and automated integration test suite).",
            "Backend Lead W.M.P.R. Weerasinghe and developers validated code via automated integration tests (tests/integration.test.cjs), verified contextIsolation security, and built the NSIS Windows installer."
        )
    ]

    for idx, (tool, purpose, outputs, verification) in enumerate(data):
        row = table.add_row()
        bg = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
        for i, val in enumerate([tool, purpose, outputs, verification]):
            c = row.cells[i]
            c.width = widths[i]
            set_cell_background(c, bg)
            set_cell_margins(c, top=90, bottom=90, left=120, right=120)
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(8.5)
            if i == 0:
                r.font.bold = True
                r.font.color.rgb = primary_color
            else:
                r.font.color.rgb = dark_text

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # Explanation callout
    callout_tbl = doc.add_table(rows=1, cols=1)
    callout_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    callout_tbl.autofit = False
    callout_tbl.columns[0].width = Inches(7.1)
    c_cell = callout_tbl.cell(0, 0)
    set_cell_background(c_cell, "EFF6FF")
    set_cell_margins(c_cell, top=100, bottom=100, left=140, right=140)

    tcPr = c_cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="2563EB"/>
            <w:top w:val="none"/>
            <w:right w:val="none"/>
            <w:bottom w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)

    cp = c_cell.paragraphs[0]
    r_c1 = cp.add_run("Summary of Correction: ")
    r_c1.font.name = "Calibri"
    r_c1.font.bold = True
    r_c1.font.size = Pt(9.5)
    r_c1.font.color.rgb = primary_color

    r_c2 = cp.add_run("This table accurately reflects reality: Gemini was used for brainstorming and research; Claude for wireframing (as evidenced by Original_Wireframes.html) and MoSCoW prioritization; Microsoft Copilot for document writing, grammar, and alignment; and Antigravity for developing the actual working prototype and codebase.")
    r_c2.font.name = "Calibri"
    r_c2.font.size = Pt(9.5)
    r_c2.font.color.rgb = dark_text

    out1 = "D:/Projects/Omega Service/AI_Tools_Usage_Table_Correction.docx"
    doc.save(out1)
    print(f"Saved: {out1}")

if __name__ == "__main__":
    generate_ai_tools_docx()
