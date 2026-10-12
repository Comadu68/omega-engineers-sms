import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="none"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="none"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def create_document():
    doc = Document()

    # Set page margins (0.8 inches)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Style colors
    primary_color = RGBColor(30, 58, 138)      # Navy #1E3A8A
    secondary_color = RGBColor(71, 85, 105)   # Slate #475569
    dark_text = RGBColor(15, 23, 42)          # Dark Slate #0F172A
    accent_red = RGBColor(185, 28, 28)        # Red #B91C1C
    accent_green = RGBColor(21, 128, 61)      # Green #15803D

    # Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("Project Architecture & Directory Tree Audit")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = primary_color

    # Subtitle
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(16)
    run_sub = sub_p.add_run("Verification Analysis & Corrected Specification for Omega Engineers SSMS")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = secondary_color
    run_sub.font.italic = True

    # Callout Box / Summary Note
    callout_tbl = doc.add_table(rows=1, cols=1)
    callout_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    callout_tbl.autofit = False
    callout_tbl.columns[0].width = Inches(6.9)
    callout_cell = callout_tbl.cell(0, 0)
    set_cell_background(callout_cell, "EFF6FF") # Light Blue
    set_cell_margins(callout_cell, top=140, bottom=140, left=180, right=180)
    
    # Left border for callout
    tcPr = callout_cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="2563EB"/>
            <w:top w:val="none"/>
            <w:right w:val="none"/>
            <w:bottom w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)

    cp = callout_cell.paragraphs[0]
    cp.paragraph_format.space_before = Pt(2)
    cp.paragraph_format.space_after = Pt(2)
    r1 = cp.add_run("Audit Finding: ")
    r1.font.name = "Calibri"
    r1.font.bold = True
    r1.font.size = Pt(10.5)
    r1.font.color.rgb = primary_color

    r2 = cp.add_run("The directory tree shown in Appendix D of the report contains several critical architectural inaccuracies when compared to the active implementation in the repository. Specifically, it violates Electron security conventions (misplacing contextBridge), misplaces backend service layers, invents non-existent filenames, misstates database schema configuration, and uses incorrect file extensions (.js vs .cjs/.jsx). Below is the comprehensive audit and the drop-in corrected replacement.")
    r2.font.name = "Calibri"
    r2.font.size = Pt(10)
    r2.font.color.rgb = dark_text

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section 1: Detailed Discrepancy Breakdown
    h1 = doc.add_heading(level=1)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    r_h1 = h1.add_run("1. Technical Discrepancy Breakdown")
    r_h1.font.name = "Calibri"
    r_h1.font.size = Pt(14)
    r_h1.font.bold = True
    r_h1.font.color.rgb = primary_color

    # Comparison Table
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="CBD5E1")

    col_widths = [Inches(1.5), Inches(2.3), Inches(3.1)]
    for i, w in enumerate(col_widths):
        table.columns[i].width = w

    # Header Row
    hdr_cells = table.rows[0].cells
    headers = ["Architecture Component", "Shown in Report Screenshot", "Actual Implementation & Reason for Error"]
    for i, h_text in enumerate(headers):
        hdr_cells[i].width = col_widths[i]
        set_cell_background(hdr_cells[i], "1E3A8A")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(h_text)
        run.font.name = "Calibri"
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    diff_data = [
        (
            "Electron Preload &\ncontextBridge\n(CRITICAL)",
            "Annotated inside /main/:\n'# Electron Main Process (OS lifecycle, window, contextBridge)'",
            "Actual: src/preload/preload.cjs\n\nWhy it's wrong: In modern Electron with contextIsolation: true, contextBridge CANNOT run in the main process. It must run inside a preload script to safely expose APIs (window.omega) to the renderer. The document completely omitted the src/preload/ folder."
        ),
        (
            "Service Layer Placement",
            "Placed at root of src/:\n/src/services/",
            "Actual: src/main/services/\n\nWhy it's wrong: All business services run inside the Node.js Main process (direct database/OS access). They belong inside src/main/services/."
        ),
        (
            "Service File Names",
            "cryptoService.js\nbillingEngine.js\nstockValidator.js\nprinterService.js",
            "Actual: 13 Domain-driven services (.cjs):\n• authService.cjs (PBKDF2 salt & hash)\n• invoiceService.cjs (billing & calculation)\n• inventoryService.cjs (stock, batches, FIFO)\n• backupService.cjs, jobService.cjs, etc.\n\nNone of the four screenshot filenames exist."
        ),
        (
            "IPC Handlers Architecture",
            "/main/index.js\n'# Main entry point & IPC registration handlers'",
            "Actual: src/main/main.cjs + src/main/ipc/\n\nWhy it's wrong: Main entry point is main.cjs (per package.json). IPC handlers are modularized into 13 separate channel modules under src/main/ipc/ rather than a single monolithic index.js."
        ),
        (
            "Database Schema &\nPersistence Layer",
            "/database/db.js\n/database/migrations/\n'# Normalized DDL scripts (12 relational tables)'",
            "Actual: src/database/schema.sql + db.cjs\n\nWhy it's wrong: No migrations/ directory exists. The canonical DDL is schema.sql, which defines 19 normalized tables and indexes (not 12), plus seed.sql."
        ),
        (
            "Renderer View Layer",
            "/renderer/screens/\n/renderer/print-layouts/",
            "Actual: src/renderer/pages/\n\nWhy it's wrong: Directory is named pages/, containing 13 React views. Printing is managed via InvoicePrintView.jsx and global.css print styles; there is no separate print-layouts/ directory."
        ),
        (
            "File Extensions",
            "Generic .js everywhere",
            "Actual: .cjs for backend, .jsx for frontend\n\nWhy it's wrong: CommonJS is required for the Electron backend runtime (.cjs), while React Vite components use .jsx."
        ),
        (
            "Repository / Root Naming",
            "Root folder: omega-engineers-ssms/\nRepo URL: omega-engineers-sms",
            "Actual Repo: omega-engineers-sms\nPackage Name: omega-engineers-service-station\n\nThere is a spelling mismatch in the document header."
        )
    ]

    for idx, (comp, doc_val, actual_val) in enumerate(diff_data):
        row = table.add_row()
        bg_col = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
        for i, val in enumerate([comp, doc_val, actual_val]):
            cell = row.cells[i]
            cell.width = col_widths[i]
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, top=100, bottom=100, left=130, right=130)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(val)
            run.font.name = "Calibri"
            run.font.size = Pt(8.5)
            if i == 0:
                run.font.bold = True
                run.font.color.rgb = primary_color
            elif i == 1:
                run.font.color.rgb = accent_red
            else:
                run.font.color.rgb = dark_text

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 2: Ready to Use Appendix D Replacement
    h2 = doc.add_heading(level=1)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)
    r_h2 = h2.add_run("2. Recommended Drop-in Replacement for Appendix D")
    r_h2.font.name = "Calibri"
    r_h2.font.size = Pt(14)
    r_h2.font.bold = True
    r_h2.font.color.rgb = primary_color

    intro_p = doc.add_paragraph()
    intro_p.paragraph_format.space_after = Pt(6)
    r_in = intro_p.add_run("You can directly copy and paste the following text into your project documentation or final report to accurately reflect the true codebase structure:")
    r_in.font.name = "Calibri"
    r_in.font.size = Pt(10)
    r_in.font.italic = True

    # Pre-formatted Box for the Tree
    tree_text = """omega-engineers-sms/
|-- /src/
|   |-- /main/                      # Electron Main Process (OS lifecycle, window orchestration)
|   |   |-- main.cjs                # Main entry point & browser window initialization
|   |   |-- /ipc/                   # Modular IPC channel registration handlers
|   |   |   |-- index.cjs           # Aggregator for all IPC registrations
|   |   |   |-- auth.ipc.cjs        # Authentication & user management IPC
|   |   |   |-- jobs.ipc.cjs        # Job card workflow & task execution IPC
|   |   |   |-- inventory.ipc.cjs   # Inventory stock, batches, & FIFO tracking IPC
|   |   |   |-- invoices.ipc.cjs    # Billing calculation & invoice generation IPC
|   |   |   |-- backup.ipc.cjs      # Database backup & restore IPC (.bak)
|   |   |   `-- ...                 # appointments, customers, mechanics, reports, etc.
|   |   `-- /services/              # Domain-Driven Business Logic & Security Services
|   |       |-- authService.cjs     # PBKDF2 hashing, per-user salt generation, RBAC
|   |       |-- jobService.cjs      # Job card state machine & mechanic assignments
|   |       |-- inventoryService.cjs# Stock thresholds, FIFO batch deductions, valuations
|   |       |-- invoiceService.cjs  # Parts + Labor + Tax billing engine & balance tracking
|   |       |-- backupService.cjs   # Automated SQLite database backup & validation
|   |       `-- ...                 # customerService, mechanicService, reportService, etc.
|   |
|   |-- /preload/                   # Secure IPC Context Isolation Layer
|   |   `-- preload.cjs             # contextBridge.exposeInMainWorld('omega', ...) API
|   |
|   |-- /database/                  # Embedded SQLite Persistence Layer
|   |   |-- db.cjs                  # better-sqlite3 connection with WAL mode & foreign keys
|   |   |-- schema.sql              # Normalized DDL schema (19 relational tables & indexes)
|   |   `-- seed.sql                # Default initial seed data & presets
|   |
|   `-- /renderer/                  # Presentation Layer (React.js + Vite UI)
|       |-- App.jsx                 # Application root & authentication router
|       |-- main.jsx                # React DOM entry point
|       |-- /components/            # Reusable UI tokens (Sidebar, Topbar, Modal, Badge)
|       |-- /pages/                 # Role-based screens & functional views
|       |   |-- Dashboard.jsx       # Operational metrics & status overview
|       |   |-- JobCards.jsx        # Job card management & dispatching
|       |   |-- Inventory.jsx       # Stock tracking, batch intake, supplier management
|       |   |-- Invoices.jsx        # Invoicing, payment collection, receipts
|       |   |-- InvoicePrintView.jsx# Print layout for A4 invoices & thermal print
|       |   |-- Reports.jsx         # Financial, job performance, and inventory reports
|       |   `-- ...                 # Appointments, Customers, Mechanics, Settings, Login
|       `-- /styles/                # Design system & CSS styling
|           `-- global.css          # Color tokens, typography, and @media print rules
|
`-- /tests/                         # Verification & Quality Assurance
    `-- integration.test.cjs        # End-to-end integration & database regression tests"""

    tree_tbl = doc.add_table(rows=1, cols=1)
    tree_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tree_tbl.autofit = False
    tree_tbl.columns[0].width = Inches(6.9)
    tree_cell = tree_tbl.cell(0, 0)
    set_cell_background(tree_cell, "0F172A") # Dark Slate / Terminal background
    set_cell_margins(tree_cell, top=140, bottom=140, left=180, right=180)

    # Border
    tcPr = tree_cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="12" w:space="0" w:color="334155"/>
            <w:top w:val="single" w:sz="12" w:space="0" w:color="334155"/>
            <w:right w:val="single" w:sz="12" w:space="0" w:color="334155"/>
            <w:bottom w:val="single" w:sz="12" w:space="0" w:color="334155"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)

    tp = tree_cell.paragraphs[0]
    tp.paragraph_format.space_before = Pt(0)
    tp.paragraph_format.space_after = Pt(0)
    tp.paragraph_format.line_spacing = Pt(11)

    r_tree = tp.add_run(tree_text)
    r_tree.font.name = "Consolas"
    r_tree.font.size = Pt(8)
    r_tree.font.color.rgb = RGBColor(226, 232, 240) # Slate 200

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 3: Summary of Defense / Viva Points
    h3 = doc.add_heading(level=1)
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(6)
    r_h3 = h3.add_run("3. Defense & Evaluation Talking Points")
    r_h3.font.name = "Calibri"
    r_h3.font.size = Pt(14)
    r_h3.font.bold = True
    r_h3.font.color.rgb = primary_color

    bullets = [
        ("Security Isolation Architecture: ", "Explain that contextBridge is explicitly implemented inside preload.cjs (rather than the main process). This enforces Electron's least-privilege security principle by ensuring the renderer process never has direct Node.js or filesystem access."),
        ("Modular Service Architecture: ", "Highlight that instead of arbitrary helper scripts (like cryptoService.js or stockValidator.js), the backend follows Domain-Driven Design (DDD) with 13 distinct services matching the business domains (Auth, Jobs, Invoices, Inventory, etc.)."),
        ("IPC Decoupling: ", "Point out that IPC handlers are partitioned across dedicated channel files in src/main/ipc/ rather than lumped into a single monolithic index file, significantly improving maintainability and testability."),
        ("Normalized Data Layer: ", "State that the SQLite schema contains 19 fully normalized tables with active foreign keys, composite indexes, and WAL (Write-Ahead Logging) mode, rather than the initial 12 tables.")
    ]

    for title_text, body_text in bullets:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_before = Pt(2)
        bp.paragraph_format.space_after = Pt(4)
        r_bt = bp.add_run(title_text)
        r_bt.font.name = "Calibri"
        r_bt.font.bold = True
        r_bt.font.size = Pt(10)
        r_bt.font.color.rgb = primary_color

        r_bb = bp.add_run(body_text)
        r_bb.font.name = "Calibri"
        r_bb.font.size = Pt(10)
        r_bb.font.color.rgb = dark_text

    # Footer note
    doc.add_paragraph().paragraph_format.space_after = Pt(16)
    foot_p = doc.add_paragraph()
    foot_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_foot = foot_p.add_run("Generated for Omega Engineers Service Station Management System • 2026")
    r_foot.font.name = "Calibri"
    r_foot.font.size = Pt(8.5)
    r_foot.font.italic = True
    r_foot.font.color.rgb = secondary_color

    # Save to file
    output_path = "D:/Projects/Omega Service/docs/Omega_Service_Architecture_Tree_Correction.docx"
    doc.save(output_path)
    print(f"Successfully generated DOCX at: {output_path}")

    # Also save a copy at root
    output_root = "D:/Projects/Omega Service/Omega_Service_Architecture_Tree_Correction.docx"
    doc.save(output_root)
    print(f"Successfully generated DOCX at: {output_root}")

if __name__ == "__main__":
    create_document()
